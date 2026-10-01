from docx import Document

file_path = "data/docx/PRFIP_Business_Data_API_Model_Documentation.docx"

document = Document(file_path)

print("Paragraphs:", len(document.paragraphs))

for i, paragraph in enumerate(document.paragraphs[:10]):
    print(f"\nParagraph {i}: {len(paragraph.text)} characters")
    print(paragraph.text[:200])